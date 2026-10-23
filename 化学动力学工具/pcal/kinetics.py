# -*- coding: utf-8 -*-
"""化学反应动力学探究（蔗糖水解）—— 数据模型与计算内核。

数据来源：Polariputer 动态测量导出的 xlsx（列：index, time, steps, angle），
或两列文本（t / α）。处理链：

    旋光度–浓度标定  α = k c + b
    α∞：实验值（反应完全后的静态测量）或由比旋光度计算的理论值
    各级反应尝试：α–t（零级）、ln(α−α∞)–t（一级）、1/(α−α∞)–t（二级）
    k1 = −slope，t1/2 = ln2/k1
    k1 = k0 + kH c(H+)ⁿ 非线性拟合 → 氢离子反应级数 n
    lg(k1−k0) = lg kH + n lg c(H+) 线性拟合作验证
"""
from __future__ import annotations

import math
import re

import numpy as np
from scipy.optimize import curve_fit
from scipy.stats import linregress, t as tdist

R = 8.314462618

# 比旋光度（20 ℃，D 线）与摩尔质量
M_SUC = 342.296
M_GLU = 180.156
M_FRU = 180.156
A_SUC = 66.6
A_GLU = 52.5
A_FRU = -91.9


# ---------------------------------------------------------------- 示例数据
def sample_data():
    """2025 年同一课程实验的完整数据（蔗糖水解，5 个盐酸浓度）。"""
    groups = [
        ('6M', 6.0, 3.015, [], [
            (131.879, 2.27748), (155.137, 2.077748), (181.226, 1.878016),
            (208.723, 1.678284), (235.204, 1.478552), (288.11, 1.143432),
            (323.607, 0.9437), (360.284, 0.743968), (407.267, 0.544236),
            (459.517, 0.344504), (518.806, 0.144772), (589.378, -0.05496),
            (677.614, -0.254692), (795.812, -0.454424), (965.083, -0.654155),
            (1345.004, -0.853887)]),
        ('5M', 5.0, 2.490, [], [
            (146.863, 2.510724), (165.113, 2.310992), (203.478, 2.11126),
            (249.067, 1.911528), (295.371, 1.711796), (345.438, 1.512064),
            (401.104, 1.312332), (464.345, 1.112601), (530.932, 0.912869),
            (605.491, 0.713137), (686.27, 0.513405), (776.663, 0.313673),
            (889.526, 0.113941), (1010.072, -0.085791), (1150.225, -0.285523),
            (1341.744, -0.485255), (1592.234, -0.684987)]),
        ('4M', 4.0, 2.000, [], [
            (128.157, 2.872654), (185.364, 2.672922), (231.041, 2.47319),
            (280.302, 2.273458), (333.815, 2.073727), (393.097, 1.873995),
            (453.561, 1.674263), (516.791, 1.474531), (585.186, 1.274799),
            (662.682, 1.075067), (748.129, 0.875335), (835.616, 0.675603),
            (946.033, 0.475871), (1071.823, 0.276139), (1213.772, 0.076408),
            (1390.677, -0.123324), (1604.966, -0.323056), (1915.627, -0.522788)]),
        ('3M', 3.0, 1.530, [], [
            (126.022, 2.969169), (201.518, 2.769437), (275.33, 2.569705),
            (347.684, 2.369973), (429.968, 2.170241), (518.456, 1.970509),
            (613.441, 1.770777), (710.579, 1.571046), (810.815, 1.371314),
            (924.522, 1.171582), (1047.007, 0.97185), (1175.64, 0.772118),
            (1334.024, 0.572386), (1505.209, 0.372654), (1689.674, 0.172922)]),
        ('2M', 2.0, 1.000, [5], [
            (140.855, 3.097855), (270.178, 2.898123), (403.39, 2.698391),
            (581.793, 2.49866), (761.599, 2.298928), (916.558, 2.099196),
            (1157.552, 1.899464), (1277.793, 1.798928), (1387.552, 1.698391),
            (1512.296, 1.597855), (1637.881, 1.497319), (1753.8, 1.396783)]),
    ]
    return {
        'temperature_c': 30.0,
        'sucrose_stock': 0.050,       # 配制的蔗糖溶液 g·mL⁻¹，混合后 1:1 稀释
        'alpha_inf_measured': -1.10,
        'calibration': [
            {'c': 0.0050, 'alpha': 0.87}, {'c': 0.0104, 'alpha': 1.51},
            {'c': 0.0250, 'alpha': 3.45}, {'c': 0.0504, 'alpha': 6.88},
            {'c': 0.0814, 'alpha': 11.11}, {'c': 0.100, 'alpha': 12.08},
            {'c': 0.120, 'alpha': 16.25}, {'c': 0.150, 'alpha': 20.60},
        ],
        'calibration_exclude': [5],
        'groups': [
            {'label': lab, 'acid_nominal': nom, 'acid_true': true,
             'alpha_inf': None, 'exclude': list(ex),
             'data': [[t, a] for t, a in pts]}
            for lab, nom, true, ex, pts in groups
        ],
    }


# ---------------------------------------------------------------- 解析
def _to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_numeric_rows(text):
    rows = []
    for line in str(text or '').replace('\r', '').split('\n'):
        line = line.strip()
        if not line or line.startswith('#'):
            continue
        parts = [p for p in re.split(r'[,;\t ]+', line) if p]
        vals = []
        ok = True
        for p in parts:
            v = _to_float(p)
            if v is None:
                ok = False
                break
            vals.append(v)
        if ok and vals:
            rows.append(vals)
    return rows


def parse_points(text):
    """把文本解析成 [[t, α], ...]。

    支持：两列 t/α；三列 index,t,α；四列 index,t,steps,angle；
    也支持带表头的文件（表头含 time / angle / alpha）。
    """
    lines = [ln.strip() for ln in str(text or '').replace('\r', '').split('\n')
             if ln.strip() and not ln.strip().startswith('#')]
    if not lines:
        return []
    header = [p.lower() for p in re.split(r'[,;\t ]+', lines[0]) if p]
    ti = ai = None
    for i, h in enumerate(header):
        if h in ('time', 't', '时间', 'time/s'):
            ti = i
        if h in ('angle', 'alpha', 'α', '旋光度', 'angle/°'):
            ai = i
    data_lines = lines[1:] if (ti is not None or ai is not None) else lines
    rows = parse_numeric_rows('\n'.join(data_lines))
    out = []
    for r in rows:
        if ti is not None and ai is not None and len(r) > max(ti, ai):
            out.append([r[ti], r[ai]])
        elif len(r) == 2:
            out.append([r[0], r[1]])
        elif len(r) == 3:
            out.append([r[1], r[2]])
        elif len(r) >= 4:
            out.append([r[1], r[-1]])
    return out


def parse_calibration(text):
    rows = parse_numeric_rows(text)
    return [{'c': r[0], 'alpha': r[1]} for r in rows if len(r) >= 2]


def parse_calibration_xlsx(source):
    """读取静态测量 xlsx：找 c/浓度 与 alpha/旋光 列；找不到就取前两列数字。"""
    import io

    import openpyxl

    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    hi = ci = ai = None
    for i, r in enumerate(rows[:10]):
        vals = [str(v).strip().lower() if v is not None else '' for v in r]
        for j, h in enumerate(vals):
            if h in ('c', 'c_suc', '浓度', 'c/(g/ml)', 'c / (g·ml-1)'):
                ci = j
            if h in ('alpha', 'α', '旋光', '旋光度', 'angle'):
                ai = j
        if ci is not None and ai is not None:
            hi = i
            break
    out = []
    if hi is not None:
        for r in rows[hi + 1:]:
            c = _to_float(r[ci]) if len(r) > ci else None
            a = _to_float(r[ai]) if len(r) > ai else None
            if c is not None and a is not None:
                out.append({'c': c, 'alpha': a})
        return out
    for r in rows:
        vals = [_to_float(v) for v in r]
        nums = [v for v in vals if v is not None]
        if len(nums) >= 2:
            out.append({'c': nums[0], 'alpha': nums[1]})
    return out


def parse_polariputer_xlsx(source):
    """读取 Polariputer 动态测量 xlsx。

    source 可以是路径或文件对象/字节流。返回
    {'points': [[t, α], ...], 'temperature': float|None, 'sucrose': float|None}。
    """
    import io

    import openpyxl

    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return {'points': [], 'temperature': None, 'sucrose': None}

    header = None
    for i, r in enumerate(rows[:10]):
        vals = [str(v).strip().lower() if v is not None else '' for v in r]
        if 'time' in vals and 'angle' in vals:
            header = (i, vals.index('time'), vals.index('angle'))
            break
    if header is None:
        # 退化成普通的两列/多列解析
        flat = '\n'.join(','.join('' if v is None else str(v) for v in r)
                         for r in rows)
        return {'points': parse_points(flat), 'temperature': None, 'sucrose': None}

    hi, ti, ai = header
    points = []
    for r in rows[hi + 1:]:
        if len(r) <= max(ti, ai):
            continue
        t, a = _to_float(r[ti]), _to_float(r[ai])
        if t is not None and a is not None:
            points.append([t, a])

    temperature = sucrose = None
    for r in rows:
        for j, v in enumerate(r):
            if not isinstance(v, str):
                continue
            s = v.strip()
            if '实验温度' in s or '温度' == s:
                for x in r[j + 1:]:
                    val = _to_float(x)
                    if val is not None:
                        temperature = val
                        break
            if '蔗糖浓度' in s:
                for x in r[j + 1:]:
                    val = _to_float(x)
                    if val is not None:
                        sucrose = val
                        break
    return {'points': points, 'temperature': temperature, 'sucrose': sucrose}


# ---------------------------------------------------------------- 拟合
def _line_fit(x, y):
    x = np.asarray(x, float)
    y = np.asarray(y, float)
    mask = np.isfinite(x) & np.isfinite(y)
    x, y = x[mask], y[mask]
    if x.size < 2:
        return None
    res = linregress(x, y)
    yhat = res.slope * x + res.intercept
    resid = y - yhat
    dof = x.size - 2
    s_err = math.sqrt(float(np.sum(resid ** 2)) / dof) if dof > 0 else float('nan')
    ci = None
    if dof > 0 and x.size >= 3:
        xbar = float(np.mean(x))
        sxx = float(np.sum((x - xbar) ** 2))
        if sxx > 0:
            tval = float(tdist.ppf(0.975, dof))
            xx = np.linspace(float(x.min()), float(x.max()), 120)
            half = tval * s_err * np.sqrt(1.0 / x.size + (xx - xbar) ** 2 / sxx)
            yy = res.slope * xx + res.intercept
            ci = dict(x=list(map(float, xx)), lo=list(map(float, yy - half)),
                      hi=list(map(float, yy + half)))
    return dict(slope=float(res.slope), intercept=float(res.intercept),
                sigma_slope=float(res.stderr or 0.0),
                sigma_intercept=float(res.intercept_stderr or 0.0),
                r2=float(res.rvalue ** 2), n=int(x.size),
                s_err=s_err, x=list(map(float, x)), y=list(map(float, y)),
                ci=ci)


def _order_y(alpha, alpha_inf, order):
    a = np.asarray(alpha, float)
    if order == 0:
        return a
    d = a - float(alpha_inf)
    if order == 1:
        y = np.full_like(d, np.nan)
        ok = d > 0
        y[ok] = np.log(d[ok])
        return y
    if order == 2:
        y = np.full_like(d, np.nan)
        ok = d != 0
        y[ok] = 1.0 / d[ok]
        return y
    raise ValueError('order must be 0/1/2')


def _fit_group(group, alpha_inf_used):
    pts = []
    for i, p in enumerate(group.get('data') or []):
        if len(p) >= 2:
            t, a = _to_float(p[0]), _to_float(p[1])
            if t is not None and a is not None:
                pts.append((i, t, a))
    exclude = set(int(i) for i in (group.get('exclude') or []))
    use = [(t, a) for i, t, a in pts if i not in exclude]
    t = np.array([p[0] for p in use], float)
    a = np.array([p[1] for p in use], float)
    fits = {}
    for order in (0, 1, 2):
        y = _order_y(a, alpha_inf_used, order)
        fits[str(order)] = _line_fit(t, y)
    first = fits['1']
    k1 = sigma_k1 = t_half = sigma_t_half = float('nan')
    alpha0_fit = float('nan')
    if first and first['slope'] < 0:
        k1 = -first['slope']
        sigma_k1 = first['sigma_slope']
        t_half = math.log(2.0) / k1
        sigma_t_half = math.log(2.0) / (k1 ** 2) * sigma_k1
        alpha0_fit = alpha_inf_used + math.exp(first['intercept'])
    return dict(label=group.get('label'), acid_nominal=group.get('acid_nominal'),
                acid_true=group.get('acid_true'), alpha_inf=alpha_inf_used,
                exclude=sorted(exclude), data=[[float(x), float(y)] for x, y in use],
                all_data=[[float(t), float(a)] for _, t, a in pts],
                n=len(use), fits=fits,
                k1=k1, sigma_k1=sigma_k1, t_half=t_half,
                sigma_t_half=sigma_t_half, alpha0_fit=alpha0_fit)


def _rate_law(cH, k):
    cH = np.asarray(cH, float)
    k = np.asarray(k, float)
    mask = np.isfinite(cH) & np.isfinite(k) & (cH > 0) & (k > 0)
    cH, k = cH[mask], k[mask]
    if cH.size < 3:
        return None

    def model(c, k0, k_h, n):
        return k0 + k_h * np.power(c, n)

    kmin, kmax = float(k.min()), float(k.max())
    p0 = [max(kmin * 0.1, 1e-7), max((kmax - kmin) / (cH.max() ** 2), 1e-8), 2.0]
    try:
        popt, pcov = curve_fit(model, cH, k, p0=p0,
                               bounds=([0.0, 0.0, 0.05], [np.inf, np.inf, 4.0]),
                               maxfev=40000)
        perr = np.sqrt(np.diag(pcov))
    except Exception:                                     # noqa: BLE001
        return None
    yhat = model(cH, *popt)
    ss_res = float(np.sum((k - yhat) ** 2))
    ss_tot = float(np.sum((k - k.mean()) ** 2))
    r2 = 1.0 - ss_res / ss_tot if ss_tot else float('nan')
    return dict(cH=list(map(float, cH)), k=list(map(float, k)),
                k0=float(popt[0]), sigma_k0=float(perr[0]),
                k_h=float(popt[1]), sigma_k_h=float(perr[1]),
                n=float(popt[2]), sigma_n=float(perr[2]), r2=float(r2),
                model='k1 = k0 + k_H c(H+)^n')


def _loglog(cH, k, k0):
    cH = np.asarray(cH, float)
    k = np.asarray(k, float)
    d = k - float(k0)
    mask = np.isfinite(cH) & np.isfinite(k) & (cH > 0) & (d > 0)
    cH, d = cH[mask], d[mask]
    if cH.size < 2:
        return None
    fit = _line_fit(np.log10(cH), np.log10(d))
    if not fit:
        return None
    fit['n'] = fit['slope']
    fit['sigma_n'] = fit['sigma_slope']
    fit['kH'] = 10.0 ** fit['intercept']
    fit['k0'] = float(k0)
    return fit


def analyze(data, cfg=None):
    cfg = dict(cfg or {})
    data = dict(data or {})
    warnings = []
    T = float(data.get('temperature_c') or 30.0)
    stock = float(data.get('sucrose_stock') or 0.0)
    c0 = stock / 2.0 if stock else float('nan')
    alpha_inf_measured = data.get('alpha_inf_measured')
    try:
        alpha_inf_measured = float(alpha_inf_measured)
    except (TypeError, ValueError):
        alpha_inf_measured = float('nan')

    # ---- 1) α–c 标定
    cal_pts = []
    for i, p in enumerate(data.get('calibration') or []):
        try:
            cal_pts.append((i, float(p['c']), float(p['alpha'])))
        except (KeyError, TypeError, ValueError):
            continue
    cal_ex = set(int(i) for i in (data.get('calibration_exclude') or []))
    use = [(c, a) for i, c, a in cal_pts if i not in cal_ex]
    cal_fit = _line_fit([p[0] for p in use], [p[1] for p in use])
    alpha0 = (cal_fit['slope'] * c0 + cal_fit['intercept']
              if cal_fit and np.isfinite(c0) else float('nan'))

    # ---- 2) α∞ 理论值
    alpha_inf_theory = float('nan')
    if np.isfinite(alpha0):
        alpha_inf_theory = alpha0 * (M_GLU / M_SUC) * (A_GLU + A_FRU) / A_SUC
    mode = cfg.get('alpha_inf_mode') or 'measured'
    alpha_inf_global = (alpha_inf_measured if mode == 'measured'
                        and np.isfinite(alpha_inf_measured)
                        else alpha_inf_theory)
    if not np.isfinite(alpha_inf_global):
        warnings.append('α∞ 没有可用值：请填实验测得的最终旋光度，或检查蔗糖浓度/标定')
    rel_ainf = ((alpha_inf_measured - alpha_inf_theory) / alpha_inf_theory * 100.0
                if np.isfinite(alpha_inf_measured)
                and np.isfinite(alpha_inf_theory) and alpha_inf_theory
                else float('nan'))

    # ---- 3) 各组动力学
    groups = []
    for g in data.get('groups') or []:
        if not g.get('label'):
            continue
        used = g.get('alpha_inf')
        try:
            used = float(used)
        except (TypeError, ValueError):
            used = alpha_inf_global
        if not np.isfinite(used):
            used = alpha_inf_global
        r = _fit_group(g, used)
        groups.append(r)
        if r['n'] < 3:
            warnings.append('%s：有效数据点只有 %d 个' % (r['label'], r['n']))
        if not np.isfinite(r['k1']):
            warnings.append('%s：一级拟合失败（检查 α−α∞ 是否为正）' % r['label'])
        elif r['fits']['1']['r2'] < 0.98:
            warnings.append('%s：一级拟合 R² = %.4f，偏低'
                            % (r['label'], r['fits']['1']['r2']))

    # ---- 4) 级数判定
    order_label = cfg.get('order_group') or (groups[1]['label'] if len(groups) > 1
                                             else (groups[0]['label'] if groups else ''))
    order_group = next((g for g in groups if g['label'] == order_label), None)
    if order_group is None and groups:
        order_group = groups[0]
    order = None
    if order_group:
        r2 = {k: (v['r2'] if v else float('nan'))
              for k, v in order_group['fits'].items()}
        best = max(r2, key=lambda k: r2[k] if np.isfinite(r2[k]) else -1)
        order = dict(group=order_group['label'], r2=r2, best=int(best),
                     is_first=(best == '1' and r2['1'] >= 0.99))

    # ---- 5) k1–c(H+) 与 lg–lg
    cH = [g['acid_true'] for g in groups]
    k1 = [g['k1'] for g in groups]
    rate_law = _rate_law(cH, k1)
    loglog = None
    if rate_law:
        loglog = _loglog(cH, k1, rate_law['k0'])
        if loglog:
            loglog['n_nonlinear'] = rate_law['n']
    if rate_law is None:
        warnings.append('k1–c(H+) 数据不足或非线性拟合失败（至少需要 3 组有效数据）')
    elif rate_law['r2'] < 0.95:
        warnings.append('k1–c(H+) 拟合 R² = %.4f，偏低；检查各组是否同一温度/同一蔗糖浓度'
                        % rate_law['r2'])

    if cal_fit is None:
        warnings.append('旋光度–浓度标定点少于 2 个，无法求 α0 与 α∞ 理论值')
    if np.isfinite(rel_ainf) and abs(rel_ainf) > 5:
        warnings.append('α∞ 实验值与理论值相差 %.1f%%，建议核对比旋光度/浓度'
                        % rel_ainf)

    return dict(
        temperature=T, c0=c0, alpha0=alpha0,
        calibration=dict(points=[{'index': i, 'c': c, 'alpha': a,
                                  'exclude': i in cal_ex}
                                 for i, c, a in cal_pts],
                         fit=cal_fit, exclude=sorted(cal_ex), n=len(use)),
        alpha_inf=dict(measured=alpha_inf_measured, theory=alpha_inf_theory,
                       used=alpha_inf_global, mode=mode, relative=rel_ainf),
        groups=groups, order=order, rate_law=rate_law, loglog=loglog,
        warnings=warnings, constants=dict(M_SUC=M_SUC, M_GLU=M_GLU,
                                          A_SUC=A_SUC, A_GLU=A_GLU, A_FRU=A_FRU))
