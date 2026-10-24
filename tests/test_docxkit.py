# -*- coding: utf-8 -*-
"""docxkit 的排版回归测试：正斜体、题注、表格、分页约束。"""
from __future__ import annotations

import os
import tempfile
import unittest

import docx

from support import REPO_ROOT, load_module, run_texts

TOOL = '电化学工具'          # 当前题注化学式规则最完整的一版
INDEX_TOOL = '双液系工具'    # ii/ij 下标规则最完整的一版


def _new_docx(tmpdir):
    path = os.path.join(tmpdir, 'blank.docx')
    docx.Document().save(path)
    return path


class DocxKitUnitTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dk = load_module(TOOL, 'docxkit')
        cls.tmp = tempfile.TemporaryDirectory()
        cls.blank = _new_docx(cls.tmp.name)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def builder(self):
        return self.dk.ReportBuilder(self.blank, cover_tail_blanks=None)

    # ---------------------------------------------------------- 题注切分
    def test_caption_variable_and_dash(self):
        self.assertEqual(self.dk._split_caption_runs('A–c'),
                         [['A', True], ['–', False], ['c', True]])

    def test_caption_chemical_species_needs_numeric_subscript(self):
        self.assertEqual(self.dk._split_caption_runs('N', '2'),
                         [['N', False]])
        self.assertEqual(self.dk._split_caption_runs('Q', 's'),
                         [['Q', True]])

    def test_caption_words_and_numbers_are_upright(self):
        for text in ('Voigt', 'A3', 'L−1', 'μmol'):
            self.assertEqual(self.dk._split_caption_runs(text),
                             [[text, False]], text)

    # ---------------------------------------------------------- 表头
    def test_header_quantity_and_unit(self):
        doc = docx.Document()
        p = doc.add_paragraph()
        self.dk.ReportBuilder._write_header(p, 'm_{Ni-1}/g', 10.5)
        runs = run_texts(p)
        self.assertEqual(runs[0][0], 'm')
        self.assertTrue(runs[0][1])
        for text in ('Ni', '-', '1', '/g'):
            hit = [r for r in runs if r[0] == text]
            self.assertTrue(hit, (text, runs))
            self.assertFalse(hit[0][1], (text, runs))

    def test_header_without_unit_still_italicizes_quantity(self):
        doc = docx.Document()
        p = doc.add_paragraph()
        self.dk.ReportBuilder._write_header(p, 'n_{0}', 10.5)
        runs = run_texts(p)
        self.assertEqual(runs[0], ('n', True, False, False, 10.5))
        self.assertEqual(runs[1][0], '0')
        self.assertFalse(runs[1][1])

    def test_chinese_header_is_upright(self):
        doc = docx.Document()
        p = doc.add_paragraph()
        self.dk.ReportBuilder._write_header(p, '轮次', 10.5)
        self.assertEqual(run_texts(p), [('轮次', False, False, False, 10.5)])

    # ---------------------------------------------------------- 题注
    def test_caption_runs_n2_and_cv(self):
        rb = self.builder()
        p = rb.caption('图3 N_{2} 的 CV 曲线')
        runs = run_texts(p)
        self.assertEqual(runs[0][0], '图3')
        self.assertTrue(runs[0][1] is False)  # 加粗在 bold，不在 italic
        self.assertTrue(p.runs[0].bold)
        n = [r for r in runs if r[0].strip() == 'N'][0]
        self.assertFalse(n[1])
        two = [r for r in runs if r[0] == '2'][0]
        self.assertTrue(two[3])                 # 真下标
        self.assertFalse(two[1])
        cv = [r for r in runs if 'CV' in r[0]]
        self.assertTrue(cv, runs)
        self.assertFalse(cv[0][1])

    def test_caption_variable_subscript_should_be_italic(self):
        # Q_{s} 的 s 是变量，必须斜体；N_{2} 的 2 是化学式下标，正体。
        rb = self.builder()
        p = rb.caption('图4 Q_{s} 与 A–c')
        s = [r for r in p.runs if r.text == 's'][0]
        self.assertTrue(s.italic)

    def test_caption_has_keep_next(self):
        rb = self.builder()
        p = rb.caption('表1 测试')
        self.assertIn('w:keepNext', p._element.xml)

    # ---------------------------------------------------------- 标题/正文
    def test_heading_has_keep_next_and_size(self):
        rb = self.builder()
        p = rb.heading('1 引言', 1)
        self.assertIn('w:keepNext', p._element.xml)
        self.assertTrue(p.runs)
        self.assertEqual(p.runs[0].font.size.pt, 10.5)

    def test_paragraph_keep_next(self):
        rb = self.builder()
        p = rb.para('直线拟合结果为：', keep_next=True)
        self.assertIn('w:keepNext', p._element.xml)

    def test_equation_is_native_omml(self):
        rb = self.builder()
        p = rb.equation(r'\frac{a}{b}')
        self.assertIn('<m:oMathPara', p._element.xml)
        self.assertIn('<m:jc m:val="center"/>', p._element.xml)
        self.assertEqual(p.alignment, 1)        # WD_ALIGN_PARAGRAPH.CENTER

    # ---------------------------------------------------------- 表格
    def test_table_structure(self):
        rb = self.builder()
        t = rb.table(['轮次', 'm_{Ni-1}/g', 'Q_{V}/J'],
                     [['1', '0.2', '100'], ['2', '0.3', '150']],
                     widths_cm=[2, 3, 3])
        self.assertEqual(len(t.rows), 3)
        self.assertEqual(len(t.columns), 3)
        self.assertIn('w:tblHeader', t.rows[0]._tr.xml)
        for row in t.rows:
            self.assertIn('w:cantSplit', row._tr.xml)
        self.assertEqual(t.cell(1, 0).text, '1')
        self.assertEqual(t.cell(2, 2).text, '150')

    def test_table_rejects_short_row(self):
        # 电导率报告里真实发生过“7 列表头、6 个数据格”导致整体左移；
        # table() 现在直接报错，而不是留下一个空末列。
        rb = self.builder()
        with self.assertRaises(ValueError):
            rb.table(['a', 'b', 'c'], [['1', '2']])

    def test_table_keep_together(self):
        rb = self.builder()
        t = rb.table(['a', 'b'], [['1', '2'], ['3', '4']],
                     keep_together=True)
        for row in t.rows[:-1]:
            for cell in row.cells:
                self.assertIn('w:keepNext', cell.paragraphs[0]._element.xml)
        for cell in t.rows[-1].cells:
            self.assertNotIn('w:keepNext', cell.paragraphs[0]._element.xml)

    # ---------------------------------------------------------- 封面
    def test_cover_long_title_gets_height_and_center(self):
        from docx.oxml import parse_xml
        from docx.oxml.ns import qn
        from support import W_NS

        doc = docx.Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '题目'
        table.cell(1, 0).text = '姓名'
        trpr = table.rows[0]._tr.get_or_add_trPr()
        trpr.append(parse_xml('<w:trHeight xmlns:w="%s" w:val="840"/>' % W_NS))
        tcpr = table.cell(0, 1)._tc.get_or_add_tcPr()
        tcpr.append(parse_xml('<w:vAlign xmlns:w="%s" w:val="top"/>' % W_NS))
        path = os.path.join(self.tmp.name, 'cover.docx')
        doc.save(path)

        rb = self.dk.ReportBuilder(path, cover_tail_blanks=None)
        matched = rb.set_cover({'题目': '稀溶液法测定极性分子的偶极矩实验报告',
                                '姓名': '测试同学'}, title_size=22.0)
        self.assertIn('题目', matched)
        table = rb.doc.tables[0]
        title_run = table.cell(0, 1).paragraphs[0].runs[0]
        self.assertEqual(title_run.font.size.pt, 22.0)
        self.assertFalse(title_run.bold)
        name_run = table.cell(1, 1).paragraphs[0].runs[0]
        self.assertEqual(name_run.font.size.pt, 14.0)
        self.assertTrue(name_run.bold)
        trh = table.rows[0]._tr.trPr.find(qn('w:trHeight'))
        self.assertEqual(trh.get(qn('w:val')), '1100')
        self.assertEqual(trh.get(qn('w:hRule')), 'atLeast')
        va = table.cell(0, 1)._tc.tcPr.find(qn('w:vAlign'))
        self.assertEqual(va.get(qn('w:val')), 'center')

    # ---------------------------------------------------------- 图片
    def test_figure_keeps_with_caption_but_caption_does_not(self):
        image = REPO_ROOT / '_tools' / 'pcal_icon.png'
        self.assertTrue(image.exists(), image)
        rb = self.builder()
        rb.figure(str(image), '图1 测试装置', width_cm=4)
        fig_p, cap_p = rb.doc.paragraphs[-2:]
        self.assertIn('w:keepNext', fig_p._element.xml)
        self.assertNotIn('w:keepNext', cap_p._element.xml)


class DocxKitIndexTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dk = load_module(INDEX_TOOL, 'docxkit')

    def test_repeated_indices_are_italic(self):
        self.assertEqual(self.dk._sub_chunks('ij'), [('ij', False)])
        self.assertEqual(self.dk._sub_chunks('ii'), [('ii', False)])

    def test_descriptive_subscripts_are_upright(self):
        self.assertEqual(self.dk._sub_chunks('sam'), [('sam', True)])
        self.assertEqual(self.dk._sub_chunks('Ni-1'),
                         [('Ni', True), ('-', True), ('1', True)])


if __name__ == '__main__':
    unittest.main()
