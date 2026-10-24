# -*- coding: utf-8 -*-
"""生成后的 docx 结构审计。

这是给“导出一份报告后先自检”用的，不依赖具体实验模块：
  · 表格列数是否与表头一致（真实发生过 7 列表头、6 个数据格）
  · 图号/表号是否从 1 开始、连续且唯一
  · 正文 run 是否五号（10.5 pt），标题是否有 keepNext
  · 是否残留 LaTeX 反斜杠（OMML 未识别的命令）
  · 是否残留 Unicode 上下标（应写成真上下标）

用法：
    python tests/audit_docx.py 报告.docx [更多.docx]
"""
from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

M_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
UNICODE_SCRIPTS = set('⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺₀₁₂₃₄₅₆₇₈₉₋₊')


@dataclass
class Issue:
    level: str          # ERROR / WARN
    code: str
    message: str
    where: str = ''

    def __str__(self):
        tail = (' @ ' + self.where) if self.where else ''
        return '[%s] %s: %s%s' % (self.level, self.code, self.message, tail)


def _is_heading(p):
    name = (p.style.name or '').lower() if p.style is not None else ''
    return 'heading' in name or '标题' in name


def _caption_number(text):
    import re
    m = re.match(r'^([图表])\s*(\d+)', (text or '').strip())
    if not m:
        return None
    return m.group(1), int(m.group(2))


def _norm_label(text):
    s = (text or '').strip()
    for ch in (' ', '\u3000', ':', '：', '\t'):
        s = s.replace(ch, '')
    return s


COVER_LABELS = ('姓名', '学号', '组别', '组号', '实验日期', '室温',
                '大气压强', '题目', '实验题目')


def _is_cover_table(table):
    """封面的「题目 / 姓名 / 学号…」信息表不是数据表，不参与列数检查。"""
    labels = []
    for row in table.rows:
        if len(row.cells) >= 2:
            labels.append(_norm_label(row.cells[0].text))
    return any(key in label for label in labels for key in COVER_LABELS)


def _check_tables(doc, issues):
    for ti, table in enumerate(doc.tables, 1):
        if not table.rows:
            issues.append(Issue('ERROR', 'table.empty', '表格 %d 没有行' % ti))
            continue
        if _is_cover_table(table):
            continue
        header = table.rows[0]
        n_head = len(header.cells)
        if n_head != len(table.columns):
            issues.append(Issue(
                'ERROR', 'table.grid',
                '表格 %d 表头 %d 格，但网格是 %d 列' % (ti, n_head, len(table.columns)),
                '表%d' % ti))
        for ri, row in enumerate(table.rows[1:], 2):
            if len(row.cells) != n_head:
                issues.append(Issue(
                    'ERROR', 'table.cols',
                    '表格 %d 第 %d 行有 %d 格，表头是 %d 格'
                    % (ti, ri, len(row.cells), n_head),
                    '表%d 第%d行' % (ti, ri)))
                continue
            if n_head and not row.cells[-1].text.strip() \
                    and header.cells[-1].text.strip():
                issues.append(Issue(
                    'WARN', 'table.trailing_empty',
                    '表格 %d 第 %d 行末列为空（表头是「%s」），'
                    '常见原因是数据行少写了一列'
                    % (ti, ri, header.cells[-1].text.strip()),
                    '表%d 第%d行' % (ti, ri)))


def _check_captions(doc, issues):
    seen = {'图': [], '表': []}
    for p in doc.paragraphs:
        text = (p.text or '').strip()
        pair = _caption_number(text)
        if not pair:
            continue
        kind, num = pair
        centered = p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        if not centered:
            # 正文里也常有「表3的……如图3所示。」这种引用。只有短、无句末
            # 标点的段落才提示“像题注但没居中”，避免每条正文都报一遍。
            short = len(text) <= 40 and not text.endswith(('。', '；', '，', '：', '！', '？'))
            if short:
                issues.append(Issue(
                    'WARN', 'caption.not_centered',
                    '「%s」看起来像题注但没有居中，可能是正文引用' % text[:30],
                    text[:30]))
            continue
        seen[kind].append((num, text))
    for kind, items in seen.items():
        if not items:
            continue
        nums = [n for n, _ in items]
        if nums[0] != 1:
            issues.append(Issue(
                'ERROR', 'caption.start',
                '%s编号从 %d 开始，应为 1' % (kind, nums[0])))
        for prev, cur in zip(nums, nums[1:]):
            if cur == prev:
                issues.append(Issue(
                    'ERROR', 'caption.duplicate',
                    '%s编号重复：%d' % (kind, cur)))
            elif cur != prev + 1:
                issues.append(Issue(
                    'ERROR', 'caption.gap',
                    '%s编号不连续：%d → %d' % (kind, prev, cur)))


def _check_text(doc, issues):
    for node in list(doc.element.body.iter(qn('w:t'))) + \
            list(doc.element.body.iter(M_NS + 't')):
        text = node.text or ''
        if '\\' in text:
            issues.append(Issue(
                'ERROR', 'text.backslash',
                '文本里残留反斜杠「%s」，通常是 OMML 未识别的命令'
                % text.strip()[:40]))
        bad = sorted(set(text) & UNICODE_SCRIPTS)
        if bad:
            issues.append(Issue(
                'WARN', 'text.unicode_script',
                '文本里残留 Unicode 上下标「%s」，应写成真上下标：%s'
                % (''.join(bad), text.strip()[:40])))


def _check_body_fonts(doc, issues):
    for p in doc.paragraphs:
        if not (p.text or '').strip():
            continue
        # 封面大标题「物理化学实验报告」是 36 pt，不属于正文；
        # 封面各字段在表格里，doc.paragraphs 本来就不会遍历到。
        if (p.text or '').strip().startswith('物理化学实验报告'):
            continue
        if _is_heading(p):
            if 'w:keepNext' not in p._element.xml:
                issues.append(Issue(
                    'WARN', 'heading.keep_next',
                    '标题没有 keepNext，可能孤零零留在页底：%s' % p.text[:30]))
        for run in p.runs:
            if not (run.text or '').strip():
                continue
            size = run.font.size
            if size is None:
                continue
            if abs(size.pt - 10.5) > 0.01:
                issues.append(Issue(
                    'WARN', 'text.font_size',
                    '正文「%s」字号是 %.1f pt，不是五号 10.5 pt'
                    % (run.text.strip()[:20], size.pt)))
                break


def audit_docx(path):
    doc = docx.Document(str(path))
    issues = []
    _check_tables(doc, issues)
    _check_captions(doc, issues)
    _check_text(doc, issues)
    _check_body_fonts(doc, issues)
    return issues


def main(argv=None):
    parser = argparse.ArgumentParser(description='审计生成的物化实验报告 docx')
    parser.add_argument('files', nargs='+')
    parser.add_argument('--strict', action='store_true',
                        help='有 WARN 也返回非零退出码')
    args = parser.parse_args(argv)

    failed = False
    for name in args.files:
        path = Path(name)
        if not path.exists():
            print('文件不存在：%s' % path)
            failed = True
            continue
        issues = audit_docx(path)
        print('\n== %s' % path)
        if not issues:
            print('  未发现问题')
            continue
        for issue in issues:
            print('  ' + str(issue))
        if any(i.level == 'ERROR' for i in issues):
            failed = True
        elif args.strict and issues:
            failed = True
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(main())
