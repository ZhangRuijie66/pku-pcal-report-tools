# -*- coding: utf-8 -*-
"""把一次 Analysis 的结果写成与“物理化学实验报告”模板同款的 .docx。"""
from __future__ import annotations

import os
import re

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import omml

FONT_HEI = '黑体'
FONT_SONG = '宋体'
FONT_FANG = '仿宋'
FONT_LATIN = 'Times New Roman'
EMU_PER_CM = 360000

# 各级标题的字体（跟课程 2026 模板一致，都是五号）：
#   一级 黑体加粗 / 二级 Arial+宋体加粗 / 三级 Arial+仿宋不加粗
HEAD_FONTS = {
    1: {'cjk': FONT_HEI, 'latin': FONT_HEI, 'bold': True},
    2: {'cjk': FONT_SONG, 'latin': 'Arial', 'bold': True},
    3: {'cjk': FONT_FANG, 'latin': 'Arial', 'bold': False},
}

W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
M_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


# ---------------------------------------------------------------- 基础排版
def style_run(run, size=10.5, bold=False, cjk=FONT_SONG, latin=FONT_LATIN,
              italic=False, color=None, sub=False, sup=False):
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if sub:
        run.font.subscript = True
    if sup:
        run.font.superscript = True
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = parse_xml('<w:rFonts %s/>' % W_NS)
        rpr.insert(0, rf)
    rf.set(qn('w:ascii'), latin)
    rf.set(qn('w:hAnsi'), latin)
    rf.set(qn('w:eastAsia'), cjk)
    rf.set(qn('w:cs'), latin)
    if cjk == FONT_HEI:
        rf.set(qn('w:hint'), 'eastAsia')


def _para_fmt(p, indent_chars=2.0, space_before=0, space_after=0,
              line_spacing=1.0, align=None, left_chars=0.0,
              page_break_before=False):
    """默认单倍行距、首行缩进 2 字符——跟课程模板的正文一致。"""
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if align is not None:
        p.alignment = align
    ppr = p._element.get_or_add_pPr()
    if page_break_before:
        # 段前分页：w:pageBreakBefore 按 schema 必须排在 spacing/ind/jc 之前
        ppr.insert(0, parse_xml('<w:pageBreakBefore %s/>' % W_NS))
    ind = ppr.find(qn('w:ind'))
    if ind is None:
        ind = parse_xml('<w:ind %s/>' % W_NS)
        ppr.append(ind)
    if indent_chars:
        ind.set(qn('w:firstLineChars'), str(int(indent_chars * 100)))
        ind.set(qn('w:firstLine'), str(int(indent_chars * 105)))
    if left_chars:
        ind.set(qn('w:leftChars'), str(int(left_chars * 100)))
        ind.set(qn('w:left'), str(int(left_chars * 105)))


# ---------------------------------------------------------------- 富文本
_RICH_RE = re.compile(r'([_^])\{([^{}]*)\}')

# Unicode 上下标字符 -> 普通字符。正文里 "J·g⁻¹"、"10⁻³"、"C₁₂H₂₂O₁₁" 这种写法
# 原来是一个固定的 Unicode 字形，字号不可控、和真正的上下标也不一致；
# 统一改写成 '^{-1}' / '_{12}' 标记，交给下面的真上下标逻辑（字号跟随正文五号）。
_UNI_SUP = {'⁰': '0', '¹': '1', '²': '2', '³': '3', '⁴': '4', '⁵': '5',
            '⁶': '6', '⁷': '7', '⁸': '8', '⁹': '9', '⁻': '-', '⁺': '+'}
_UNI_SUB = {'₀': '0', '₁': '1', '₂': '2', '₃': '3', '₄': '4', '₅': '5',
            '₆': '6', '₇': '7', '₈': '8', '₉': '9', '₋': '-', '₊': '+'}

# 符号位上的连字符（负号/减号）：前面不是字母数字或右括号，后面紧跟数字
_MINUS_RE = re.compile(r'(?<![\w)])-(?=\d)')


def _minus_signs(text):
    """
    把"负号/减号"的 ASCII 连字符换成真减号 U+2212。

    只改真正的符号位：前面不是字母数字或右括号、后面紧跟数字。
    这样 '−26424'、'^{-1}'、'10^{-3}' 会换，而 'm_{Ni-1}'、'5-75'、'GPL-3.0' 不会。
    """
    return _MINUS_RE.sub('\u2212', text)


def _scripts_to_markers(text):
    """把 Unicode 上下标改写成 ^{..} / _{..} 标记。"""
    if not any((ch in _UNI_SUP) or (ch in _UNI_SUB) for ch in text):
        return text
    out = []
    i, n = 0, len(text)
    while i < n:
        if text[i] in _UNI_SUP:
            table, mark = _UNI_SUP, '^'
        elif text[i] in _UNI_SUB:
            table, mark = _UNI_SUB, '_'
        else:
            out.append(text[i])
            i += 1
            continue
        buf = []
        while i < n and text[i] in table:
            buf.append(table[text[i]])
            i += 1
        out.append('%s{%s}' % (mark, ''.join(buf)))
    return ''.join(out)


# 正斜体判定（普通文本里的"变量 run"）：
#   · 拉丁字母、小写希腊字母 = 变量 → 斜体
#   · 大写希腊字母（Δ、Σ…）、数字、运算符、空格 = 正体
# 所以 'Δ T' 要写成两个 run：正体的 Δ + 斜体的 T。
_GREEK_LOWER = set('αβγδεζηθικλμνξοπρστυφχψω')
_GREEK_UPPER = set('ΑΒΓΔΕΖΗΘΙΚΛΜΝΞΟΠΡΣΤΥΦΧΨΩ')


def _is_var_char(ch):
    """斜体变量字符：拉丁字母 + 小写希腊字母。"""
    return ('a' <= ch <= 'z') or ('A' <= ch <= 'Z') or (ch in _GREEK_LOWER)


def _split_runs_by_italic(text):
    """把一段变量文本按正/斜体切成 [(子串, italic), …]。"""
    out = []
    for ch in text:
        it = _is_var_char(ch)
        if out and out[-1][1] is it:
            out[-1][0] += ch
        else:
            out.append([ch, it])
    return out


def _split_caption_runs(text, next_body=None):
    """
    图题/表题里只把**量符号**排斜体，普通单词与编号保持正体：
        'A3'      → 正体 A3（分子编号）
        'Voigt'   → 正体（人名）
        'μmol'    → 正体（单位，μ 是词头）
        'L−1'     → 正体 L（单位升）
        'A–c'     → 斜体 A、正体 –、斜体 c
        'A 类'    → 正体 A（类别标签）

    next_body：紧随其后的上下标内容（题注切分时的「下标预览」）——
    单字母 + 纯数字下标（N_{2}、O_{2}）是化学物种，字母正体；
    单字母 + 字母下标（Q_{s}）才是变量，斜体。
    """
    out = []
    i, n = 0, len(text)
    while i < n:
        ch = text[i]
        if ('a' <= ch <= 'z') or ('A' <= ch <= 'Z'):
            j = i
            while j < n and (('a' <= text[j] <= 'z') or ('A' <= text[j] <= 'Z')):
                j += 1
            word, nxt = text[i:j], (text[j] if j < n else '')
            if len(word) > 1:
                it = False                      # Voigt、mol、nm…
            else:
                it = True
                if nxt.isdigit():               # A3、B2：编号正体
                    it = False
                elif word == 'L' and nxt in '·−':   # L·mol⁻¹ / L−1：单位升
                    it = False
                elif nxt == ' ' and text[j + 1:j + 2] == '类':   # A 类
                    it = False
                elif j >= n and next_body is not None:
                    body = next_body.strip()
                    if not body or not any(c.isalpha() for c in body):
                        it = False              # N_{2}：化学式下标纯数字
            out.append([word, it])
            i = j
            continue
        if ch in _GREEK_LOWER:
            prev_letter = i > 0 and text[i - 1].isascii() and text[i - 1].isalpha()
            nxt_letter = i + 1 < n and text[i + 1].isascii() and text[i + 1].isalpha()
            it = not (ch == 'μ' and nxt_letter and not prev_letter)  # μmol/μm 正体
            out.append([ch, it])
            i += 1
            continue
        out.append([ch, False])
        i += 1
    # 合并相邻同样式的片段
    merged = []
    for chunk, it in out:
        if merged and merged[-1][1] is it:
            merged[-1][0] += chunk
        else:
            merged.append([chunk, it])
    return merged


_SUB_CHUNK_RE = re.compile(r'[A-Za-z]+|\d+|[^A-Za-z\d]+')


def _sub_chunks(body):
    """
    下标的逐块正斜体：多字母（元素符号、sam/cot/tot 这类缩写）与数字、
    连字符一律正体；单字母是变量本身，斜体。
    'V-Ni' → 斜体的 V + 正体的 '-' + 正体的 Ni
    """
    out = []
    for m in _SUB_CHUNK_RE.finditer(body):
        s = m.group(0)
        up = (not s.isalpha()) or len(s) > 1
        out.append((s, up))
    return out


def _emit_text(p, text, opts):
    """按 opts 写一段文本；斜体 run 里的大写希腊字母 / 数字 / 符号自动改正体。"""
    if not text:
        return
    if (opts or {}).get('caption_italic'):
        for chunk, italic in _split_caption_runs(text):
            o = {k: v for k, v in opts.items() if k != 'caption_italic'}
            o['italic'] = italic
            style_run(p.add_run(chunk), **o)
        return
    if not (opts or {}).get('italic'):
        style_run(p.add_run(text), **(opts or {}))
        return
    for chunk, italic in _split_runs_by_italic(text):
        o = dict(opts)
        o['italic'] = italic
        style_run(p.add_run(chunk), **o)


def _write_script(p, body, opts, sup):
    """
    上下标逐块写：说明性文字（sam、Ni…）与数字、符号一律正体，
    单字母变量（V、p、n…）在变量 run 里保持斜体。
    所以 Q_{V-Ni} → 斜体 V + 正体 -Ni；R^{2} → 斜体 R + 正体上标 2。
    """
    for chunk, upright in _sub_chunks(body):
        o = {k: v for k, v in opts.items() if k != 'caption_italic'}
        if o.get('italic'):
            o['italic'] = not upright
        r = p.add_run(chunk)
        if sup:
            r.font.superscript = True
        else:
            r.font.subscript = True
        style_run(r, **o)


def _rich_tokens(text):
    """把含 _{..} / ^{..} 标记的文本切成 [(kind, body), …]，kind: t/sub/sup。"""
    out, pos = [], 0
    for m in _RICH_RE.finditer(text):
        if m.start() > pos:
            out.append(('t', text[pos:m.start()]))
        out.append(('sub' if m.group(1) == '_' else 'sup', m.group(2)))
        pos = m.end()
    if pos < len(text):
        out.append(('t', text[pos:]))
    return out


def add_rich_text(p, text, base_opts=None, sub_scale=1.0, caption=False):
    """
    把 'm_{Ni-1}/g'、'10^{-5}' 这类写法写成真正的 Word 下标 / 上标。

    用于表格表头、正文里的化学式等场合——这些地方只是记录物理量的名字，
    不应该（也不需要）用公式对象。

    sub_scale=1.0：上下标的字号属性与正文同为五号（10.5 pt）。Word 显示上下标
    时本来就会自己缩小并升降，这里再乘一次系数会让它们小得看不清。
    """
    opts = dict(base_opts or {})
    text = _minus_signs(_scripts_to_markers(str(text)))
    if caption:
        # 题注：正斜体切分要看得到下标内容（N_{2} 是化学式 → 正体，Q_{s} 是变量 → 斜体）
        toks = _rich_tokens(text)
        for idx, (kind, body) in enumerate(toks):
            if kind != 't':
                sub = dict(opts)
                sub['italic'] = False
                if sub_scale != 1.0:
                    sub['size'] = round(sub.get('size', 10.5) * sub_scale, 1)
                _write_script(p, body, sub, sup=(kind == 'sup'))
                continue
            nxt = None
            if idx + 1 < len(toks) and toks[idx + 1][0] != 't':
                nxt = toks[idx + 1][1]
            for chunk, italic in _split_caption_runs(body, next_body=nxt):
                o = dict(opts)
                o['italic'] = italic
                style_run(p.add_run(chunk), **o)
        return
    pos = 0
    for m in _RICH_RE.finditer(text):
        if m.start() > pos:
            _emit_text(p, text[pos:m.start()], opts)
        sub = dict(opts)
        if sub_scale != 1.0:
            sub['size'] = round(sub.get('size', 10.5) * sub_scale, 1)
        _write_script(p, m.group(2), sub, sup=(m.group(1) == '^'))
        pos = m.end()
    if pos < len(text):
        _emit_text(p, text[pos:], opts)


def _add_text(p, text, opts=None):
    """写一段正文：含 _{..} / ^{..} 时用真上下标，否则走单个 run（快）。"""
    text = _minus_signs(_scripts_to_markers('' if text is None else str(text)))
    if _RICH_RE.search(text):
        add_rich_text(p, text, opts)
    else:
        _emit_text(p, text, opts or {})


class ReportBuilder:
    def __init__(self, template_path, cover_tail_blanks=1):
        self.doc = docx.Document(template_path)
        self.template_dir = os.path.dirname(os.path.abspath(template_path))
        self._head_styles = {}
        if cover_tail_blanks is not None:
            self._trim_cover_blanks(cover_tail_blanks)

    # ------------------------------------------------------------ 样式
    def _heading_style(self, level):
        """
        取模板里的「标题 1/2/3」样式。用样式（而不是直接排版）是课程要求的：
        导出 PDF 时才会带上书签目录。模板里这几个样式的 styleId 是 1/2/3。

        模板只有三级标题；四级（4.3.1 这种）沿用「标题 3」的样式，
        字体由 HEAD_FONTS.get(level, HEAD_FONTS[3]) 兜底，
        这样报告里可以出现 4.3.1，而 PDF 书签最多到三级。
        """
        if level in self._head_styles:
            return self._head_styles[level]
        names = {1: ('heading 1', '标题 1'), 2: ('heading 2', '标题 2'),
                 3: ('heading 3', '标题 3')}.get(level, ('heading 3', '标题 3'))
        want_id = str(min(level, 3))
        found = None
        for st in self.doc.styles:
            try:
                nm = (st.name or '').strip().lower()
            except Exception:                                  # noqa: BLE001
                continue
            if nm in names:
                found = st
                break
            if found is None and getattr(st, 'style_id', None) == want_id:
                found = st            # 名称对不上就用 styleId 兜底
        self._head_styles[level] = found
        return found

    def _trim_cover_blanks(self, keep):
        """
        删掉封面信息表之后、正文之前多余的空行，好让「摘要」「关键词」留在第一页。

        模板封面结构：标题 + 题目表 + 空行 + 信息表 + N 行空行。
        从末尾往前数，跳过纯装饰段（带 drawing 的空段），把多余的空行删掉。
        """
        body = self.doc.element.body
        paras = body.findall(qn('w:p'))
        if not paras:
            return
        blanks = []
        for p in reversed(paras):
            has_text = bool(''.join(t.text or '' for t in p.iter(qn('w:t'))).strip())
            if has_text:
                break
            if p.findall('.//' + qn('w:drawing')):
                continue                    # 装饰段留着，不算空行
            blanks.append(p)
        for p in blanks[:max(0, len(blanks) - keep)]:
            body.remove(p)

    # ------------------------------------------------------------ 封面
    @staticmethod
    def _norm_label(text):
        s = (text or '').strip()
        for ch in (' ', '\u3000', ':', '：', '\t'):
            s = s.replace(ch, '')
        return s

    def set_cover(self, mapping, title_size=26.0):
        """
        按“字段名 -> 值”更新封面表格（只覆盖右侧单元格）。返回命中的字段。

        字号跟课程模板的段落标记一致：题目默认 26 pt 黑体；其余字段 14 pt 黑体加粗。
        题目过长（例如本实验的 25 字题目）时用 title_size 调小，否则会撑破单元格。
        """
        matched = []
        for table in self.doc.tables:
            for row in table.rows:
                cells = row.cells
                if len(cells) < 2:
                    continue
                label = self._norm_label(cells[0].text)
                for key, value in mapping.items():
                    if label == self._norm_label(key) and value not in (None, ''):
                        cell = cells[1]
                        for p in cell.paragraphs[1:]:
                            p._element.getparent().remove(p._element)
                        p = cell.paragraphs[0]
                        for r in list(p.runs):
                            r._element.getparent().remove(r._element)
                        big = label in ('题目', '实验题目', '题目名称')
                        size = float(title_size) if big else 14.0
                        r = p.add_run(str(value))
                        style_run(r, size=size,
                                  cjk=FONT_HEI, latin=FONT_HEI,
                                  bold=not big)
                        if big:
                            # 长题目要能折行：单倍行距，并把单元格加高、垂直居中
                            p.paragraph_format.line_spacing = 1.0
                            p.paragraph_format.space_before = 0
                            p.paragraph_format.space_after = 0
                            trpr = row._tr.get_or_add_trPr()
                            trh = trpr.find(qn('w:trHeight'))
                            if trh is not None:
                                trh.set(qn('w:val'), '1100')
                                trh.set(qn('w:hRule'), 'atLeast')
                            tcpr = cell._tc.get_or_add_tcPr()
                            va = tcpr.find(qn('w:vAlign'))
                            if va is not None:
                                va.set(qn('w:val'), 'center')
                        # 段落标记也带上同样格式，模板里单元格是空的，靠这个定字号
                        ppr = p._element.get_or_add_pPr()
                        rpr = ppr.find(qn('w:rPr'))
                        if rpr is not None:
                            ppr.remove(rpr)
                        ppr.append(parse_xml(
                            '<w:rPr %s><w:rFonts w:ascii="%s" w:eastAsia="%s" '
                            'w:hAnsi="%s"%s/><w:sz w:val="%d"/><w:szCs w:val="%d"/>'
                            '%s</w:rPr>'
                            % (W_NS, FONT_HEI, FONT_HEI, FONT_HEI,
                               ' w:hint="eastAsia"',
                               int(size * 2), int(size * 2),
                               '' if big else '<w:b/>')))
                        matched.append(key)
        return matched

    def set_header_title(self, title):
        """把页眉里的「实验题目」换成真正的题目（模板页眉是 左 | 中 | 右 的制表位布局）。"""
        for section in self.doc.sections:
            for p in section.header.paragraphs:
                for r in p.runs:
                    if '实验题目' in (r.text or ''):
                        r.text = str(title)
                        return True
        return False

    # ------------------------------------------------------------ 文本
    def heading(self, text, level=1, page_break_before=False):
        """
        text —— 字符串，或与 para() 相同的片段列表。
        标题里可能带引用上标（「2 实验部分^{[1]}」）或量符号
        （「3.2.2 量热计常数 W 的计算」→ W 用 ('t','W',ITALIC)）。

        标题**必须套模板里的「标题 1/2/3」样式**：一来字体字号跟课程模板一致，
        二来导出 PDF 才会带书签目录（课程模板里用红字专门要求过）。
        """
        p = self.doc.add_paragraph()
        st = self._heading_style(level)
        if st is not None:
            p.style = st
        _para_fmt(p, indent_chars=0, line_spacing=1.0,
                  space_before=6, space_after=6,
                  page_break_before=page_break_before)
        # 标题不能孤零零留在页底：跟后面的正文/公式绑在一页
        _ppr = p._element.get_or_add_pPr()
        if _ppr.find(qn('w:keepNext')) is None:
            _ppr.insert(0, parse_xml('<w:keepNext %s/>' % W_NS))
        # 段落标记也要压到五号，否则模板样式里的 22 pt 会把标题行高撑大
        ppr = p._element.get_or_add_pPr()
        old = ppr.find(qn('w:rPr'))
        if old is not None:
            ppr.remove(old)
        ppr.append(parse_xml('<w:rPr %s><w:sz w:val="21"/><w:szCs w:val="21"/></w:rPr>'
                             % W_NS))
        opts = {'size': 10.5}
        opts.update(HEAD_FONTS.get(level, HEAD_FONTS[3]))
        self._write_segments(p, text, opts)
        return p

    def para(self, segments, indent_chars=2.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             line_spacing=1.0, keep_next=False):
        """
        segments —— 字符串，或段落片段列表。片段可以是：
            '普通文本'                     默认正文格式
            ('t', '文本')                  带格式的文本
            ('t', '文本', {'bold': True})  指定格式（cjk/size/latin/bold/italic）
            ('m', latex)                   行内公式
            ('文本', {'bold': True})       等价写法，向后兼容
        """
        p = self.doc.add_paragraph()
        _para_fmt(p, indent_chars=indent_chars, align=align, line_spacing=line_spacing)
        if keep_next:
            ppr = p._element.get_or_add_pPr()
            if ppr.find(qn('w:keepNext')) is None:
                ppr.insert(0, parse_xml('<w:keepNext %s/>' % W_NS))
        self._write_segments(p, segments)
        return p

    def _write_segments(self, p, segments, base_opts=None):
        """把片段列表写进段落；base_opts 是整段的默认格式（标题会用）。"""
        base = dict(base_opts or {})
        if isinstance(segments, str):
            segments = [segments]
        for seg in segments:
            if isinstance(seg, str):
                _add_text(p, seg, base)
                continue
            if isinstance(seg, dict):
                opts = dict(base)
                opts.update({k: v for k, v in seg.items() if k != 'text'})
                _add_text(p, seg.get('text', ''), opts)
                continue
            if not isinstance(seg, (tuple, list)) or not seg:
                _add_text(p, str(seg), base)
                continue
            kind = seg[0]
            if kind == 'm':
                self._add_math(p, seg[1], display=False)
                continue
            if kind == 't':
                extra = seg[2] if len(seg) > 2 and isinstance(seg[2], dict) else {}
                text = seg[1]
            else:
                # 旧式 ('文本', {opts}) / ('文本', {})
                extra = seg[1] if len(seg) > 1 and isinstance(seg[1], dict) else {}
                text = seg[0]
            opts = dict(base)
            opts.update(extra)
            _add_text(p, text, opts)
        return p

    def equation(self, latex):
        """独立成行的居中公式（Word 原生公式）。"""
        p = self.doc.add_paragraph()
        _para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5)
        self._add_math(p, latex, display=True)
        return p

    def _add_math(self, p, latex, display=False):
        if display:
            xml = ('<m:oMathPara %s %s><m:oMathParaPr><m:jc m:val="center"/></m:oMathParaPr>'
                   '<m:oMath>%s</m:oMath></m:oMathPara>') % (W_NS, M_NS, omml.to_omml_body(latex))
        else:
            xml = '<m:oMath %s %s>%s</m:oMath>' % (W_NS, M_NS, omml.to_omml_body(latex))
        p._element.append(parse_xml(xml))

    def caption(self, text, keep_next=True):
        """图题 / 表题：居中，「表1」「图2」这几个字加粗（跟课程模板一致）。"""
        p = self.doc.add_paragraph()
        _para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
                  space_before=0, space_after=6)
        # keepNext：题注必须与紧随其后的表格 / 图片留在同一页，
        # 否则表题会孤零零地留在上一页页底。
        if keep_next:
            ppr = p._element.get_or_add_pPr()
            if ppr.find(qn('w:keepNext')) is None:
                ppr.insert(0, parse_xml('<w:keepNext %s/>' % W_NS))
        head = ''
        rest = text or ''
        m = re.match(r'^([表图]\s*\d+)(.*)$', rest)
        if m:
            head, rest = m.group(1), m.group(2)
        if head:
            style_run(p.add_run(head), size=10.5, bold=True)
        # 题注里可能出现量符号（Q_s、A–c…）：真上/下标 + 只把量符号排斜体
        # （A3、Voigt、μmol、L 这类编号/单词/单位保持正体）
        add_rich_text(p, rest, {'size': 10.5}, caption=True)
        return p

    def blank(self, size=10.5):
        p = self.doc.add_paragraph()
        _para_fmt(p, indent_chars=0, line_spacing=1.0)
        r = p.add_run('')
        style_run(r, size=size)
        return p

    # ------------------------------------------------------------ 图片
    def figure(self, image_path, caption=None, width_cm=14.0):
        if not os.path.exists(image_path):
            return
        p = self.doc.add_paragraph()
        _para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0)
        if caption:
            ppr = p._element.get_or_add_pPr()      # 图片与图题不能分家
            if ppr.find(qn('w:keepNext')) is None:
                ppr.insert(0, parse_xml('<w:keepNext %s/>' % W_NS))
        r = p.add_run()
        r.add_picture(image_path, width=Cm(width_cm))
        if caption:
            self.caption(caption, keep_next=False)

    # ------------------------------------------------------------ 表格
    def table(self, headers, rows, widths_cm=None, font_size=10.5,
              cell_margin_mm=0.9, keep_together=False):
        """表格默认用五号字（10.5 pt）；不自动换行，靠给定列宽排开。"""
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = self._grid_style()
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        self._set_cell_margins(t, cell_margin_mm)
        tblpr = t._tbl.tblPr
        layout = tblpr.find(qn('w:tblLayout'))
        if layout is None:
            layout = parse_xml('<w:tblLayout %s w:type="fixed"/>' % W_NS)
            tblpr.append(layout)
        for j, h in enumerate(headers):
            cell = t.cell(0, j)
            cell.text = ''
            p = cell.paragraphs[0]
            _para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0)
            # 表头同样是普通文本（支持 m_{Ni-1} 这类真下标），不用公式对象
            self._write_header(p, h, font_size)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                cell = t.cell(i + 1, j)
                cell.text = ''
                p = cell.paragraphs[0]
                _para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0)
                self._write_cell(p, v, font_size)
        if widths_cm:
            for j, w in enumerate(widths_cm):
                for row in t.rows:
                    row.cells[j].width = Cm(w)
        self._row_props(t)
        if keep_together:
            # 小表不要跨页：除最后一行外，每行的段落都加 keepNext
            for row in t.rows[:-1]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        ppr = p._element.get_or_add_pPr()
                        if ppr.find(qn('w:keepNext')) is None:
                            ppr.insert(0, parse_xml('<w:keepNext %s/>' % W_NS))
        return t

    @staticmethod
    def _row_props(table):
        """
        表格跨页时的两条规矩：
          · 表头行加 w:tblHeader，续页自动重复表头；
          · 每一行加 w:cantSplit，绝不把一行从中间劈开
            （否则「第3轮」会被断成「第3」/「轮」分处两页）。
        """
        for i, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            tags = ['w:cantSplit'] + (['w:tblHeader'] if i == 0 else [])
            for tag in tags:
                if trpr.find(qn(tag)) is None:
                    trpr.append(parse_xml('<%s %s/>' % (tag, W_NS)))

    @staticmethod
    def _write_header(p, text, font_size):
        """
        表头：'m_{Ni-1}/g' 里量符号是变量（斜体），'/' 之后的单位是正体。
        中文表头（轮次、用途）不受影响。

        ⚠️ 没有单位的表头（如 'n_{0}'，一列纯量符号）也必须按量符号排斜体 ——
        早先只在含 '/' 时才对斜杠前那段做斜体，于是 'n_{0}' 走了 else 分支
        整块正体，表 1 / 表 3 的 n₀ 列头就成了粗正体（用户 2026-09-23 抓到）。
        判据：表头里出现拉丁字母或小写希腊字母 → 是量符号，斜体；纯中文才正体。
        """
        text = '' if text is None else str(text)
        base = {'size': font_size, 'bold': True}
        if '/' in text:
            sym, unit = text.split('/', 1)
            add_rich_text(p, sym, dict(base, italic=True))
            add_rich_text(p, '/' + unit, base)
        elif re.search(r'[A-Za-zα-ω]', text):
            add_rich_text(p, text, dict(base, italic=True))
        else:
            add_rich_text(p, text, base)

    @staticmethod
    def _set_cell_margins(table, mm):
        """收紧单元格内边距，避免因为默认边距把文字挤到换行。"""
        twips = int(round(mm * 56.7))
        tblpr = table._tbl.tblPr
        mar = tblpr.find(qn('w:tblCellMar'))
        if mar is not None:
            tblpr.remove(mar)
        mar = parse_xml(
            '<w:tblCellMar %s>'
            '<w:top w:w="%d" w:type="dxa"/><w:left w:w="%d" w:type="dxa"/>'
            '<w:bottom w:w="%d" w:type="dxa"/><w:right w:w="%d" w:type="dxa"/>'
            '</w:tblCellMar>' % (W_NS, twips, twips, twips, twips))
        tblpr.append(mar)

    def _write_cell(self, p, value, font_size):
        if isinstance(value, tuple) and value and value[0] == 'm':
            self._add_math(p, value[1])
            return
        # 纯文本单元格：支持 m_{Ni-1}/g 这种真下标写法，不再使用公式对象
        add_rich_text(p, '' if value is None else str(value), {'size': font_size})

    def _grid_style(self):
        return 'Table Grid'

    # ------------------------------------------------------------ 保存
    def save(self, path):
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.doc.save(path)
        return path
